# Weeks, dates and hosts from the scoreboard sheet's "Schedule" tab (Week | Start Date | Host).
# Writes weeks.json ([week, "YYYY-MM-DD", host, holidays]) and data_week.txt (latest completed week).
import openpyxl, json, datetime, re, os, sys

today = datetime.date.fromisoformat(os.environ["BUILD_DATE"]) if os.environ.get("BUILD_DATE") else datetime.date.today()
hist = {w[0]: w for w in json.load(open("weeks_hist.json"))}
LOST = "[LOST DATA]"

def to_date(v):
    if isinstance(v, datetime.datetime): return v.date()
    if isinstance(v, datetime.date): return v
    s = str(v or "").strip()
    for fmt in ("%B %d, %Y", "%b %d, %Y", "%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y"):
        try: return datetime.datetime.strptime(s, fmt).date()
        except ValueError: pass
    return None

wb = openpyxl.load_workbook("board.xlsx", read_only=True, data_only=True)
rows = {}
if "Schedule" in wb.sheetnames:
    for r in wb["Schedule"].iter_rows(min_row=2, values_only=True):
        if not r or r[0] in (None, ""): continue
        try: w = int(float(r[0]))
        except (TypeError, ValueError): continue
        d = to_date(r[1] if len(r) > 1 else None)
        host = str(r[2]).strip() if len(r) > 2 and r[2] not in (None,) else ""
        if host.upper() == "[LOST DATA]": host = LOST
        rows[w] = (d, host)
    print(f"Schedule tab: {len(rows)} weeks")
else:
    print("WARNING: no Schedule tab found; using the saved host history only", file=sys.stderr)

start = datetime.date(2013, 10, 26)  # week 1
last = max(list(rows) + list(hist))
weeks = []
for w in range(1, last + 1):
    d, host = rows.get(w, (None, None))
    if d is None: d = start + datetime.timedelta(weeks=w - 1)
    if host is None: host = hist[w][2] if w in hist else ""
    # anniversary marker: the week containing October 22 (Rocksmith 2014 / Championship birthday)
    ann = ""
    if d.year >= 2014 and d <= datetime.date(d.year, 10, 22) <= d + datetime.timedelta(days=6):
        ann = f"RS2014 & Champ {d.year - 2013} Year Anniversary"
    weeks.append([w, d.isoformat(), host, ann])  # other holidays are worked out in the page from the calendar

# latest completed week: started at least 7 days ago, and the next week has begun. A week only ends once
# BOTH the next week's WeekNNN tab exists on the scoreboard sheet AND it's past 8 pm Central European time (Europe/Berlin: CET in winter, CEST in summer) on the
# Saturday the next week starts. The real clock is used (BUILD_DATE only sets the date the workflows build for;
# a BUILD_DATE other than today, for testing, is treated as the end of that day).
UTC = datetime.timezone.utc
now = datetime.datetime.now(UTC)
from zoneinfo import ZoneInfo
if os.environ.get("BUILD_DATE") and today != now.astimezone(ZoneInfo("America/New_York")).date():
    now = datetime.datetime.combine(today, datetime.time(23, 59), UTC)
tabs = {int(m.group(1)) for s in wb.sheetnames for m in [re.match(r"Week(\d+)$", s)] if m}
starts = {w: datetime.date.fromisoformat(d) for w, d, h, a in weeks}
def next_started(w):
    n = w + 1
    if n not in starts: return False
    if tabs and n >= min(tabs) and n not in tabs: return False   # scoreboard era: the tab must exist
    return now >= datetime.datetime.combine(starts[n], datetime.time(20, 0), ZoneInfo("Europe/Berlin"))   # 8 pm CET/CEST
done = [w for w, d, h, a in weeks if datetime.date.fromisoformat(d) + datetime.timedelta(days=7) <= today and next_started(w)]
data_week = max(done)
json.dump(weeks, open("weeks.json", "w"), separators=(",", ":"))
open("data_week.txt", "w").write(str(data_week))
print(f"weeks 1-{last}, latest completed week {data_week}")
