import openpyxl,re,json
wb=openpyxl.load_workbook('board.xlsx',data_only=True)
out={}
for ws in wb.worksheets:
    m=re.fullmatch(r'Week(\d+)',ws.title)
    if not m or ws.max_column!=136 or ws.cell(8,2).value!='Name': continue
    wk=int(m.group(1)); chal=[]
    maxr=min(ws.max_row,120)
    for b in range(27):
        c=2+5*b
        song=ws.cell(6,c).value; lvl=ws.cell(7,c).value; diff=ws.cell(7,c+3).value; extra=ws.cell(7,c+2).value
        ents=[]
        for r in range(9,maxr+1):
            n=ws.cell(r,c).value
            if n in (None,''): continue
            vals=[ws.cell(r,c+k).value for k in (1,2,3)]
            ents.append([str(n).strip()]+[v if isinstance(v,(int,float)) else None for v in vals])
        if (song in (None,'','n/a')) and not ents: continue
        chal.append([b//9,b%9,str(lvl or ''),str(song or ''),diff if isinstance(diff,(int,float)) else None,str(extra or ''),ents])
    out[wk]=chal
json.dump(out,open('scores_raw.json','w'),separators=(',',':'))
print('weeks',len(out),min(out),max(out))
