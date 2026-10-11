# Read the WeekNNN scoreboard tabs: values (scores_raw.json) and name fonts (fonts.json, used to tell
# who was playing down). MODE=full reads every tab; MODE=quick only reads the newest weeks and reuses
# the saved results (cache/) for the rest.
import openpyxl, re, json, os

MODE = os.environ.get("MODE", "full")
QUICK_WEEKS = int(os.environ.get("QUICK_WEEKS", "2"))   # how many of the newest week tabs a quick run re-reads
wb = openpyxl.load_workbook("board.xlsx", read_only=True, data_only=True)
tabs = {int(m.group(1)): n for n in wb.sheetnames if (m := re.fullmatch(r"Week(\d+)", n))}

raw, fonts, skipped, arr = {}, {}, set(), {}
if MODE == "quick" and os.path.exists("cache/scores_raw.json"):
    raw = json.load(open("cache/scores_raw.json"))
    fonts = json.load(open("cache/fonts.json"))
    if os.path.exists("cache/skipped.json"): skipped = set(json.load(open("cache/skipped.json")))
    if os.path.exists("cache/arr.json"): arr = json.load(open("cache/arr.json"))
    newest = sorted(tabs)[-QUICK_WEEKS:]
    todo = [w for w in tabs if w in newest or (str(w) not in raw and w not in skipped)]
else:
    MODE = "full"; todo = list(tabs)

def color(f):
    c = f.color
    if c is None: return None
    return c.rgb if c.type == "rgb" else "theme"

for wk in sorted(todo):
    ws = wb[tabs[wk]]
    grid = [list(r) for r in ws.iter_rows(min_row=1, max_row=120, max_col=136)]
    if len(grid) < 8: skipped.add(wk); continue
    cell = lambda r, c: grid[r - 1][c - 1] if r - 1 < len(grid) and c - 1 < len(grid[r - 1]) else None
    val = lambda r, c: getattr(cell(r, c), "value", None)
    # current scoreboard layout: 27 blocks of 5 columns (Lead/Rhythm/Bass x 9 levels), "Name" headers on row 8
    if val(8, 2) != "Name" or val(8, 132) != "Name": skipped.add(wk); continue   # older layout
    chal, frows, sw = [], {}, {}
    for b in range(27):
        c = 2 + 5 * b
        song, lvl, diff, extra = val(6, c), val(7, c), val(7, c + 3), val(7, c + 2)
        # the code after the level is the song arrangement to play: L/R/B (normally the path's own, but some weeks swap
        # it, e.g. a Lead challenge played on the Rhythm arrangement), BL/BR/BB (a bonus arrangement inside the song)
        # or AL/AR/AB (an alternate arrangement)
        al = re.sub(r"[^A-Z]", "", str(val(7, c + 1) or "").upper())
        if al in ("L", "R", "B", "BL", "BR", "BB", "AL", "AR", "AB") and al != "LRB"[b // 9]: sw[str(b)] = al
        ents, fl = [], []
        for r in range(9, len(grid) + 1):
            n = val(r, c)
            if n in (None, ""): continue
            vals = [val(r, c + k) for k in (1, 2, 3)]
            ents.append([str(n).strip()] + [v if isinstance(v, (int, float)) else None for v in vals])
            f = cell(r, c).font
            fl.append([str(n).strip(), color(f), bool(f.b), f.sz])
        frows[b] = fl
        if song in (None, "", "n/a") and not ents: continue
        chal.append([b // 9, b % 9, str(lvl or ""), str(song or ""), diff if isinstance(diff, (int, float)) else None, str(extra or ""), ents])
    raw[str(wk)] = chal
    if sw: arr[str(wk)] = sw
    else: arr.pop(str(wk), None)
    if wk >= 440: fonts[str(wk)] = frows

# keep the sheet's tab order, and drop weeks whose tab was deleted
order = [str(w) for w in tabs]
raw = {k: raw[k] for k in order if k in raw}
fonts = {k: fonts[k] for k in order if k in fonts}
os.makedirs("cache", exist_ok=True)
arr = {k: arr[k] for k in order if k in arr}
for name, obj in (("scores_raw.json", raw), ("fonts.json", fonts), ("arr.json", arr)):
    s = json.dumps(obj, separators=(",", ":"))
    open(name, "w").write(s); open("cache/" + name, "w").write(s)
json.dump(sorted(skipped), open("cache/skipped.json", "w"))
print(f"{MODE} scan: read {len(todo)} week tabs, {len(raw)} weeks total")
