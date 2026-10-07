# Rodman's Rankings tab: the leaderboard of the last 10 weeks, per path and level -> rankings.json
import openpyxl, json, re

wb = openpyxl.load_workbook("board.xlsx", read_only=True, data_only=True)
name = next((n for n in wb.sheetnames if re.fullmatch(r"Rodman.s Rankings", n.strip())), None)
out = {"rules": "", "weeks": "", "updated": "", "listed": [], "paths": []}
if name:
    rows = [list(r) for r in wb[name].iter_rows(min_row=1, max_row=400, max_col=24, values_only=True)]
    cell = lambda r, c: rows[r][c] if r < len(rows) and c < len(rows[r]) else None
    s = lambda v: "" if v is None else (str(int(v)) if isinstance(v, float) and v.is_integer() else str(v)).strip()
    out["rules"] = s(cell(0, 8)); out["weeks"] = s(cell(2, 3)); out["updated"] = s(cell(3, 3))
    out["listed"] = [s(cell(2, c)) for c in range(8, 12)]
    BLOCKS = [1, 6, 11, 16]                      # Rank column of each level block (Beginner, Intermediate, Advanced, Masterclass)
    i = 0
    while i < len(rows):
        t = s(cell(i, 1))
        m = re.search(r"(Lead|Rhythm|Bass)", t) if t.startswith("~") else None
        if not m: i += 1; continue
        path = m.group(1); levels = None; i += 1
        while i < len(rows) and s(cell(i, 1)) != "Rank":
            if s(cell(i, 1)) == "Beginner": levels = [s(cell(i, c)) for c in BLOCKS]
            i += 1
        i += 1
        lv = [{"level": (levels or ["Beginner", "Intermediate", "Advanced", "Masterclass"])[k], "rows": []} for k in range(4)]
        while i < len(rows):
            got = False
            for k, c in enumerate(BLOCKS):
                rank, nm, pts, chg = s(cell(i, c)), s(cell(i, c + 1)), s(cell(i, c + 2)), s(cell(i, c + 3))
                if nm: lv[k]["rows"].append([rank, nm, pts, chg]); got = True
            if not got: break
            i += 1
        out["paths"].append({"path": path, "levels": lv})
json.dump(out, open("rankings.json", "w"), ensure_ascii=False, separators=(",", ":"))
print("rankings:", out["weeks"], [(p["path"], [len(l["rows"]) for l in p["levels"]]) for p in out["paths"]])
