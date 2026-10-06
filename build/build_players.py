# Player list, attendance and players-per-week.
# Attendance history (old Metrics workbook) is saved for weeks 2-675. From week 440 on, anyone on that week's
# scoreboard also counts as played (the scoreboard fills in players the Metrics workbook missed); from week 676 on
# the scoreboard is the only source. Week 565's scoreboard tab is incomplete, so history is kept rather than replaced.
import openpyxl, json, collections, sys

ALIAS = {'AndreCardoso': 'ALSRC', 'BetterCallJamie': 'FlatCap Jay', 'jasmith85': 'FlatCap Jay', 'aand14': 'drand',
         'mattimatti': 'mattibassi', 'Tinman86': 'TinMan86'}
DROP = {'#REF!', 'Player Name'}
SB_START = 440
data_week = int(open("data_week.txt").read())
H = json.load(open("presence_hist.json"))
hist, HIST_END = H["players"], H["through_week"]     # name -> [tiers, bits for weeks 2..HIST_END]
parts_hist = json.load(open("parts_hist.json"))     # players per week, weeks 2..439
sb = json.load(open("scores.json"))

# PlayersList tab: name, Lead, Rhythm, Bass, ...
wb = openpyxl.load_workbook("board.xlsx", read_only=True, data_only=True)
tiers = {}
for r in wb["PlayersList"].iter_rows(min_row=2, values_only=True):
    if not r or not r[0]: continue
    n = str(r[0]).strip()
    if n in ALIAS: continue
    t = "".join((str(x).strip().upper()[:1] if x not in (None, "") else "?") for x in r[1:4])
    tiers[n] = t

canon = {}
for n in hist: canon.setdefault(n.lower(), n)
for n in tiers: canon[n.lower()] = n                 # the PlayersList spelling wins

n_hist = min(HIST_END, data_week) - 1
nbits = data_week - 1                                # weeks 2..data_week
pres = collections.defaultdict(lambda: bytearray(b"0" * nbits))
for n, (t, bits) in hist.items():
    c = canon[n.lower()]
    b = pres[c]
    for i, ch in enumerate(bits[:n_hist]):
        if ch == "1": b[i] = ord("1")

unknown = collections.Counter()
for w in range(SB_START, data_week + 1):
    seen = set()
    for card in sb["weeks"].get(str(w), []):
        for e in card[6]:
            raw = sb["names"][e[0]].rstrip("*~").strip()
            raw = ALIAS.get(raw, raw)
            if raw in DROP or not raw: continue
            c = canon.get(raw.lower())
            if c is None: unknown[raw] += 1; continue
            seen.add(c)
    for c in seen: pres[c][w - 2] = ord("1")
# players per week: saved Metrics counts before the scoreboard era, then counted from attendance
parts = list(parts_hist[:SB_START - 2]) + [sum(1 for b in pres.values() if b[w - 2] == ord("1")) for w in range(SB_START, data_week + 1)]

players = []
for n, b in pres.items():
    bits = b.decode()
    ws = [i + 2 for i, ch in enumerate(bits) if ch == "1"]
    if not ws: continue
    best = run = 0; prev = None
    for w in ws:
        run = run + 1 if prev == w - 1 else 1; prev = w; best = max(best, run)
    act = 0; w = data_week
    while w >= 2 and bits[w - 2] == "1": act += 1; w -= 1
    first = ws[0]
    padded = bits + "0" * (-len(bits) % 4)
    hx = "".join(format(int(padded[i:i + 4], 2), "x") for i in range(0, len(padded), 4))
    t = tiers.get(n) or (hist.get(n) or [None])[0]
    players.append([n, first, ws[-1], len(ws), round(len(ws) / (data_week - first + 1), 4), act, best, hx, t])

json.dump(players, open("players.json", "w"), separators=(",", ":"))
json.dump(parts, open("parts.json", "w"))
print(f"players {len(players)}, weeks 2-{data_week}")
if unknown:
    print("Names on the scoreboard that aren't on PlayersList (not counted):", file=sys.stderr)
    for n, k in unknown.most_common(): print(f"  {n} ({k} entries)", file=sys.stderr)
