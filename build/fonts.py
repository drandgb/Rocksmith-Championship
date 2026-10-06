import openpyxl,re,json,collections
wb=openpyxl.load_workbook('board.xlsx')
out={}
for ws in wb.worksheets:
    m=re.fullmatch(r'Week(\d+)',ws.title)
    if not m or ws.max_column!=136 or ws.cell(8,2).value!='Name': continue
    wk=int(m.group(1))
    if wk<440: continue
    rows={}
    for b in range(27):
        c=2+5*b; L=[]
        for r in range(9,min(ws.max_row,120)+1):
            cell=ws.cell(r,c)
            if cell.value in (None,''): continue
            f=cell.font; col=f.color
            rgb=col.rgb if col is not None and col.type=='rgb' else ('theme' if col is not None else None)
            L.append([str(cell.value).strip(),rgb,bool(f.b),f.sz])
        rows[b]=L
    out[wk]=rows
json.dump(out,open('fonts.json','w'))
print(len(out))
